"""
Telegram-бот для оформления расписания пар.

Два режима работы:

1. Присылаешь боту сырой текст пары (как раньше) — он отвечает
   отформатированным сообщением с эмодзи и ссылкой на zoom.
2. Присылаешь боту .docx-файл с расписанием (таблица на несколько групп,
   как выгружает деканат) — бот сам находит нужную группу, разбирает все
   пары и отвечает готовым расписанием на день. Дату можно указать в
   подписи к файлу (например "15.09"), иначе берётся сегодняшняя дата.

Эмодзи для конкретных предметов и преподавателей, а также нужная группа —
настраиваются прямо в чате командами, без правки кода.

Запуск:
    pip install -r requirements.txt
    export BOT_TOKEN="токен_от_BotFather"      (или создай файл token.txt)
    python bot.py

Подробности — в README.md.
"""

from __future__ import annotations

import asyncio
import json
import logging
import os
import re
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

from aiogram import Bot, Dispatcher, F, Router
from aiogram.enums import ParseMode
from aiogram.exceptions import TelegramBadRequest
from aiogram.filters import Command, CommandObject, CommandStart
from aiogram.types import Message
from docx import Document

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

BASE_DIR = Path(__file__).resolve().parent
DATA_FILE = BASE_DIR / "data.json"

# ---- эмодзи по умолчанию (используются, пока для предмета/преподавателя
#      не настроен свой) ------------------------------------------------
DEFAULT_SUBJECT_EMOJI = "📕"
DEFAULT_TEACHER_EMOJI = "👤"
CLOCK_EMOJI = "🕔"
LINK_EMOJI = "💻"
ROOM_EMOJI = "🚪"

DEFAULT_GROUP = "ДХМоС-36"
KYIV_TZ = timezone(timedelta(hours=3))

# ---- ключевые слова учёных/преподавательских званий, по которым бот
#      отделяет название предмета от преподавателя внутри одной строки.
#      Отсортированы от самых длинных к самым коротким, чтобы более
#      точная фраза не перебивалась короткой ("ст. викл." раньше "викл.") --
TITLE_KEYWORDS = [
    "старший викладач", "ст. викл.", "ст.викл.",
    "завідувач кафедри", "зав. кафедри", "зав.каф.",
    "професор", "проф.",
    "доцент", "доц.",
    "асистент", "асист.", "ас.",
    "викладач", "викл.",
]
TITLE_KEYWORDS.sort(key=len, reverse=True)

# Прізвище + ініціали, напр. "Ярешко С.П." або "Ярешко С. П."
NAME_PATTERN = re.compile(
    r"[А-ЯЁІЇЄҐ][а-яёіїєґ'’ʼ\-]+\s+[А-ЯЁІЇЄҐ]\.\s?[А-ЯЁІЇЄҐ]\."
)

PAIR_START_RE = re.compile(r"^\s*(\d+)\s*пара", re.IGNORECASE | re.MULTILINE)
URL_RE = re.compile(r"https?://\S+")
TIME_RE = re.compile(r"(\d{1,2}[:.]\d{2})\s*[-–—]\s*(\d{1,2}[:.]\d{2})")

# ---- разбор .docx-таблицы расписания ----------------------------------
TYPE_MARKERS_RE = re.compile(r"(Лек\.|Лаб\.|Пр\.|Сем\.|Конс\.|МК\.|ПК\.?)", re.IGNORECASE)
DATE_TOKEN_RE = re.compile(r"\b(\d{1,2})\.(\d{1,2})\b")
ROOM_RE = re.compile(r"ауд\.?\s*([^,]*)", re.IGNORECASE)
URLSAFE_RE = re.compile(r"^[A-Za-z0-9\-_./?=&%:]+$")
DATE_IN_CAPTION_RE = re.compile(r"(\d{1,2})[.\-/](\d{1,2})")


# ---------------------------------------------------------------------
# хранение настроек (эмодзи для предметов/преподавателей, группа) в data.json
# ---------------------------------------------------------------------
def load_data() -> dict:
    if DATA_FILE.exists():
        try:
            payload = json.loads(DATA_FILE.read_text(encoding="utf-8"))
            payload.setdefault("subjects", {})
            payload.setdefault("teachers", {})
            payload.setdefault("group", DEFAULT_GROUP)
            return payload
        except json.JSONDecodeError:
            logger.warning("data.json повреждён, начинаю с чистого листа")
    return {"subjects": {}, "teachers": {}, "group": DEFAULT_GROUP}


def save_data(payload: dict) -> None:
    DATA_FILE.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )


data = load_data()


# ---------------------------------------------------------------------
# общие вспомогательные функции форматирования
# ---------------------------------------------------------------------
_DIGIT_KEYCAPS = {d: f"{d}️⃣" for d in "0123456789"}


def keycap(number: int) -> str:
    """1 -> 1️⃣, 12 -> 1️⃣2️⃣ и т.д."""
    return "".join(_DIGIT_KEYCAPS[d] for d in str(number))


def find_subject_emoji(subject_text: str) -> str:
    low = subject_text.lower()
    best: tuple[str, str] | None = None
    for key, emoji in data["subjects"].items():
        klow = key.lower()
        if (klow in low or low in klow) and (best is None or len(key) > len(best[0])):
            best = (key, emoji)
    return best[1] if best else DEFAULT_SUBJECT_EMOJI


def find_teacher_emoji(teacher_text: str) -> str:
    """
    Ищет настроенный эмодзи преподавателя. Работает и если админ задал только
    фамилию ("Ярешко"), и если задал целиком "ст. викл. Ярешко С.П." — ключ и
    искомый текст сравниваются в обе стороны, как подстроки друг друга.
    """
    low = teacher_text.lower()
    best: tuple[str, str] | None = None
    for key, emoji in data["teachers"].items():
        klow = key.lower()
        if (klow in low or low in klow) and (best is None or len(key) > len(best[0])):
            best = (key, emoji)
    return best[1] if best else DEFAULT_TEACHER_EMOJI


def split_subject_teacher(line: str) -> tuple[str, str, str]:
    """
    Делит строку вида "Рисунок і живопис ст. викл. Ярешко С.П."
    на (предмет, строка преподавателя, ключ для поиска эмодзи преподавателя).
    """
    low = line.lower()
    title_pos = None
    for kw in TITLE_KEYWORDS:
        idx = low.find(kw.lower())
        if idx != -1 and (title_pos is None or idx < title_pos):
            title_pos = idx

    if title_pos is not None:
        subject = line[:title_pos].strip(" ,.-")
        teacher_line = line[title_pos:].strip()
    else:
        m = NAME_PATTERN.search(line)
        if m:
            subject = line[: m.start()].strip(" ,.-")
            teacher_line = line[m.start():].strip()
        else:
            subject = line.strip()
            teacher_line = ""

    m = NAME_PATTERN.search(teacher_line)
    teacher_key = m.group(0).split()[0] if m else teacher_line

    return subject, teacher_line, teacher_key


# =======================================================================
# Режим 1: разбор сырого текста, который присылает пользователь вручную
# =======================================================================
def format_block(block: str) -> str | None:
    lines = [ln.strip() for ln in block.strip().splitlines() if ln.strip()]
    if not lines:
        return None

    m = PAIR_START_RE.match(lines[0])
    if not m:
        return None
    pair_num = int(m.group(1))

    url = None
    text_lines = []
    for ln in lines[1:]:
        um = URL_RE.search(ln)
        if um:
            url = um.group(0)
            continue
        text_lines.append(ln)

    time_str = None
    remaining = []
    for ln in text_lines:
        tm = TIME_RE.search(ln)
        if tm and time_str is None:
            time_str = f"{tm.group(1)}-{tm.group(2)}"
            leftover = (ln[: tm.start()] + ln[tm.end():]).strip(" ()")
            if leftover:
                remaining.append(leftover)
        else:
            remaining.append(ln)

    subject_teacher_line = " ".join(remaining).strip()
    subject, teacher_line, teacher_key = split_subject_teacher(subject_teacher_line)

    room = None
    rm = ROOM_RE.search(teacher_line)
    if rm:
        room = rm.group(1).strip(" .")
        teacher_line = (teacher_line[: rm.start()] + teacher_line[rm.end():]).strip(" ,")

    subject_emoji = find_subject_emoji(subject) if subject else DEFAULT_SUBJECT_EMOJI
    teacher_emoji = find_teacher_emoji(teacher_line) if teacher_line else DEFAULT_TEACHER_EMOJI

    out = [f"{keycap(pair_num)}пара"]
    if time_str:
        out.append(f"{CLOCK_EMOJI} {time_str}")
    if subject:
        out.append(f"{subject_emoji}{subject}")
    if teacher_line:
        out.append(f"{teacher_emoji}{teacher_line}")
    if room:
        out.append(f"{ROOM_EMOJI} ауд. {room}")
    if url:
        out.append(f"{LINK_EMOJI} [zoom]({url})")

    return "\n".join(out)


def format_schedule(raw_text: str) -> str:
    """Форматирует одну или сразу несколько пар из одного сообщения."""
    starts = list(PAIR_START_RE.finditer(raw_text))
    if not starts:
        return ""

    blocks = []
    for i, m in enumerate(starts):
        start = m.start()
        end = starts[i + 1].start() if i + 1 < len(starts) else len(raw_text)
        blocks.append(raw_text[start:end])

    formatted = [format_block(b) for b in blocks]
    formatted = [f for f in formatted if f]
    return "\n\n".join(formatted)


# =======================================================================
# Режим 2: разбор .docx-таблицы расписания (несколько групп в одной
# таблице, лекция/лаба чередуются по конкретным датам)
# =======================================================================
def _looks_urlsafe(s: str) -> bool:
    return bool(URLSAFE_RE.match(s))


def _extract_url(paragraphs: list[str]) -> tuple[str | None, list[str]]:
    """
    Достаёт ссылку из списка абзацев ячейки. Иногда Word переносит длинную
    ссылку на несколько абзацев внутри одной ячейки — склеиваем их обратно.
    Всё, что стоит в том же абзаце сразу после самой ссылки (например,
    "Ідентифікатор конференції: ..."), в ссылку не попадает.
    """
    texts = [p.strip() for p in paragraphs if p.strip()]
    url = None
    rest: list[str] = []
    i, n = 0, len(texts)
    while i < n:
        t = texts[i]
        if url is None and t.lower().startswith("http"):
            buf = t
            i += 1
            while i < n and _looks_urlsafe(texts[i]):
                buf += texts[i]
                i += 1
            um = URL_RE.search(buf)
            url = um.group(0).rstrip(".,;\"") if um else buf.rstrip(".,;")
            continue
        rest.append(t)
        i += 1
    return url, rest


# служебные пометки про идентификатор конференции / пароль — не нужны в выводе
MEETING_NOTE_RE = re.compile(
    r"ідентифікатор конференц|код доступ|meeting id|passcode", re.IGNORECASE
)


def _entry_dates(text: str) -> list[str]:
    return sorted({f"{int(dd):02d}.{int(mm):02d}" for dd, mm in DATE_TOKEN_RE.findall(text)})


def _parse_cell(cell) -> dict | None:
    paragraphs = [p.text for p in cell.paragraphs]
    url, rest = _extract_url(paragraphs)
    rest = [t for t in rest if not MEETING_NOTE_RE.search(t)]
    if not rest:
        return None

    dates = _entry_dates(" ".join(rest))

    # аудиторию ищем в пределах ОДНОГО абзаца, чтобы случайно не "утащить"
    # в неё текст из соседних, не связанных с ней абзацев
    room = ""
    for idx, t in enumerate(rest):
        m = ROOM_RE.search(t)
        if m:
            room = m.group(1).strip(" .")
            rest[idx] = (t[: m.start()] + t[m.end():]).strip(" ,")
            break

    text_wo_room = " ".join(x for x in rest if x)
    clean = TYPE_MARKERS_RE.sub("", text_wo_room)
    clean = DATE_TOKEN_RE.sub("", clean)
    clean = re.sub(r"[ ,]{2,}", " ", clean).strip(" ,")

    subject, teacher_line, teacher_key = split_subject_teacher(clean)
    if not subject and not teacher_line:
        return None

    return {
        "subject": subject,
        "teacher_line": teacher_line,
        "teacher_key": teacher_key,
        "room": room,
        "url": url,
        "dates": dates,
    }


def _find_group_columns(table, group_name: str) -> list[int]:
    if not table.rows:
        return []
    header = table.rows[0].cells
    gnorm = group_name.strip().lower().replace(" ", "")
    cols = []
    for idx, cell in enumerate(header):
        first_line = cell.text.split("\n")[0].strip().lower().replace(" ", "")
        if gnorm and first_line and (gnorm in first_line or first_line in gnorm):
            cols.append(idx)
    return cols


def extract_group_schedule(doc_bytes: bytes, group_name: str) -> dict[int, dict]:
    """
    Пара -> {"joint": [...], "subgroups": {1: [...], 2: [...]}}.

    Группа в таблице деканата иногда занимает две физические колонки —
    одну на подгруппу. Когда для конкретной пары обе колонки — это на
    самом деле одна объединённая (склеенная) ячейка Word, значит пара
    общая на всю группу ("joint"). Когда это две разные ячейки — у
    подгрупп разные предмет/преподаватель/аудитория, и их нужно
    показывать отдельно.
    """
    document = Document(BytesIO(doc_bytes))
    raw: dict[int, dict] = defaultdict(
        lambda: {"joint": [], "subgroups": defaultdict(list), "max_subgroups": 1}
    )

    for table in document.tables:
        cols = _find_group_columns(table, group_name)
        if not cols:
            continue
        for row in table.rows[1:]:
            pair_text = row.cells[0].text.strip()
            if not pair_text.isdigit():
                continue
            pair_num = int(pair_text)
            time_text = row.cells[1].text.strip()
            valid_cols = [c for c in cols if c < len(row.cells)]
            if not valid_cols:
                continue

            bucket = raw[pair_num]
            bucket["max_subgroups"] = max(bucket["max_subgroups"], len(valid_cols))

            # группируем колонки по тому, ссылаются ли они на одну и ту же
            # (склеенную) ячейку Word — это и есть признак "общая пара"
            merged_groups: list[dict] = []
            for pos, col in enumerate(valid_cols):
                cell = row.cells[col]
                match = next((g for g in merged_groups if g["tc"] is cell._tc), None)
                if match:
                    match["positions"].append(pos + 1)
                else:
                    merged_groups.append({"tc": cell._tc, "cell": cell, "positions": [pos + 1]})

            if len(merged_groups) == 1:
                parsed = _parse_cell(merged_groups[0]["cell"])
                if parsed and parsed["subject"]:
                    parsed["time"] = time_text
                    bucket["joint"].append(parsed)
            else:
                for g in merged_groups:
                    parsed = _parse_cell(g["cell"])
                    if not parsed or not parsed["subject"]:
                        continue
                    parsed["time"] = time_text
                    for subgroup_num in g["positions"]:
                        bucket["subgroups"][subgroup_num].append(parsed)

    return raw


def _dedup_candidates(candidates: list[dict]) -> list[dict]:
    uniq: list[dict] = []
    seen = set()
    for c in candidates:
        key = (c["subject"], c["teacher_line"], c["room"], c["url"], tuple(c["dates"]))
        if key not in seen:
            seen.add(key)
            uniq.append(c)
    return uniq


def _pick_single(candidates: list[dict], target_date_str: str) -> tuple[dict | None, str | None]:
    uniq = _dedup_candidates(candidates)
    if not uniq:
        return None, None

    exact = [c for c in uniq if target_date_str in c["dates"]]
    if exact:
        return exact[0], None

    # точного совпадения нет — вариант без привязки к датам считаем
    # действующим всегда
    undated = [c for c in uniq if not c["dates"]]
    if undated:
        return undated[0], None

    # все оставшиеся варианты жёстко привязаны к конкретным датам, и ни
    # один не подошёл под нужный день
    if len(uniq) == 1:
        # единственная запись про эту пару, но именно на эту дату она не
        # распространяется — скорее всего, в этот день пары просто нет
        return None, None

    # несколько дато-зависимых вариантов — берём самый вероятный (с самым
    # длинным списком дат), но помечаем как неточный
    uniq_sorted = sorted(uniq, key=lambda c: len(c["dates"]), reverse=True)
    return uniq_sorted[0], "uncertain"


def pick_for_date(
    raw_entries: dict[int, dict], target_date_str: str
) -> dict[int, list[tuple[int | None, dict, str | None]]]:
    """
    Для каждой пары выбирает вариант(ы), которые действуют на нужную дату.
    Возвращает пара -> список (номер_подгруппы_или_None, запись, пометка).
    Если у пары есть подгруппы, но на выбранную дату их содержимое совпало —
    отдаёт один общий вариант без пометки подгруппы. Если для какой-то
    подгруппы на эту дату вообще ничего не нашлось — её просто не показываем
    (а не приписываем ей чужую запись).
    """
    result: dict[int, list[tuple[int | None, dict, str | None]]] = {}
    for pair_num, bucket in raw_entries.items():
        joint = bucket["joint"]
        subgroups = bucket["subgroups"]
        max_sg = bucket.get("max_subgroups", 1)

        if max_sg <= 1:
            chosen, note = _pick_single(joint, target_date_str)
            if chosen:
                result[pair_num] = [(None, chosen, note)]
            continue

        picks: list[tuple[int, dict, str | None]] = []
        contents = set()
        for sg in range(1, max_sg + 1):
            chosen, note = _pick_single(joint + subgroups.get(sg, []), target_date_str)
            if chosen:
                picks.append((sg, chosen, note))
                contents.add((chosen["subject"], chosen["teacher_line"], chosen["room"], chosen["url"]))

        if not picks:
            continue
        if len(picks) == max_sg and len(contents) <= 1:
            # у всех подгрупп нашлось что-то, и оно одинаковое — это общая пара
            sg0, chosen0, note0 = picks[0]
            result[pair_num] = [(None, chosen0, note0)]
        else:
            result[pair_num] = picks

    return result


SUBGROUP_EMOJI = {1: "🅰️", 2: "🅱️"}
SUBGROUP_LABEL = {1: "1 підгрупа", 2: "2 підгрупа"}


def format_docx_entry(pair_num: int, entry: dict, subgroup: int | None = None) -> str:
    subject_emoji = find_subject_emoji(entry["subject"]) if entry["subject"] else DEFAULT_SUBJECT_EMOJI
    teacher_emoji = find_teacher_emoji(entry["teacher_line"]) if entry["teacher_line"] else DEFAULT_TEACHER_EMOJI

    header = f"{keycap(pair_num)}пара"
    if subgroup is not None:
        emoji = SUBGROUP_EMOJI.get(subgroup, f"{keycap(subgroup)}")
        label = SUBGROUP_LABEL.get(subgroup, f"{subgroup} підгрупа")
        header += f" {emoji} {label}"

    out = [header]
    if entry.get("time"):
        out.append(f"{CLOCK_EMOJI} {entry['time']}")
    if entry["subject"]:
        out.append(f"{subject_emoji}{entry['subject']}")
    if entry["teacher_line"]:
        out.append(f"{teacher_emoji}{entry['teacher_line']}")
    if entry["room"]:
        out.append(f"{ROOM_EMOJI} ауд. {entry['room']}")
    if entry["url"]:
        out.append(f"{LINK_EMOJI} [zoom]({entry['url']})")
    return "\n".join(out)


# ---------------------------------------------------------------------
# бот
# ---------------------------------------------------------------------
router = Router()

HELP_TEXT = (
    "Умею работать двумя способами.\n\n"
    "1️⃣ Пришли текст расписания (можно сразу за несколько пар подряд) — "
    "оформлю его с эмодзи и красивой ссылкой на zoom.\n\n"
    "2️⃣ Пришли .docx-файл с расписанием (таблица на несколько групп) — сам "
    "найду настроенную группу и оформлю все пары на день. Чтобы выбрать дату, "
    "укажи её в подписи к файлу, например «15.09» — иначе возьму сегодняшнюю. "
    "Если в этот день у подгрупп разные пары — покажу обе, с пометкой "
    "🅰️ 1 підгрупа / 🅱️ 2 підгрупа.\n\n"
    "Настройки:\n"
    "/setgroup <название группы> — какую группу доставать из .docx (сейчас: "
    f"«{data.get('group', DEFAULT_GROUP)}»)\n"
    "/setsubject <эмодзи> <название предмета> — например:\n"
    "/setsubject 🎨 Рисунок і живопис\n"
    "/setteacher <эмодзи> <прізвище> — например:\n"
    "/setteacher 👨‍🎨 Ярешко\n\n"
    "/subjects — список настроенных предметов\n"
    "/teachers — список настроенных преподавателей\n"
    "/delsubject <название> — удалить настройку предмета\n"
    "/delteacher <прізвище> — удалить настройку преподавателя"
)


@router.message(CommandStart())
async def cmd_start(message: Message) -> None:
    await message.answer(HELP_TEXT)


@router.message(Command("help"))
async def cmd_help(message: Message) -> None:
    await message.answer(HELP_TEXT)


@router.message(Command("setgroup"))
async def cmd_setgroup(message: Message, command: CommandObject) -> None:
    name = (command.args or "").strip()
    if not name:
        await message.answer(
            f"Сейчас настроена группа «{data.get('group', DEFAULT_GROUP)}». "
            "Формат: /setgroup Назва групи"
        )
        return
    data["group"] = name
    save_data(data)
    await message.answer(f"Готово, теперь достаю из .docx группу «{name}»")


@router.message(Command("setsubject"))
async def cmd_setsubject(message: Message, command: CommandObject) -> None:
    if not command.args:
        await message.answer("Формат: /setsubject <эмодзи> <название предмета>")
        return
    parts = command.args.split(maxsplit=1)
    if len(parts) < 2:
        await message.answer("Формат: /setsubject <эмодзи> <название предмета>")
        return
    emoji, name = parts[0], parts[1].strip()
    data["subjects"][name] = emoji
    save_data(data)
    await message.answer(f"Готово: «{name}» → {emoji}")


@router.message(Command("setteacher"))
async def cmd_setteacher(message: Message, command: CommandObject) -> None:
    if not command.args:
        await message.answer("Формат: /setteacher <эмодзи> <прізвище>")
        return
    parts = command.args.split(maxsplit=1)
    if len(parts) < 2:
        await message.answer("Формат: /setteacher <эмодзи> <прізвище>")
        return
    emoji, name = parts[0], parts[1].strip()
    data["teachers"][name] = emoji
    save_data(data)
    await message.answer(f"Готово: «{name}» → {emoji}")


@router.message(Command("delsubject"))
async def cmd_delsubject(message: Message, command: CommandObject) -> None:
    name = (command.args or "").strip()
    if name in data["subjects"]:
        del data["subjects"][name]
        save_data(data)
        await message.answer(f"Удалено: «{name}»")
    else:
        await message.answer("Не нашёл такой предмет. Посмотри /subjects")


@router.message(Command("delteacher"))
async def cmd_delteacher(message: Message, command: CommandObject) -> None:
    name = (command.args or "").strip()
    if name in data["teachers"]:
        del data["teachers"][name]
        save_data(data)
        await message.answer(f"Удалено: «{name}»")
    else:
        await message.answer("Не нашёл такого преподавателя. Посмотри /teachers")


@router.message(Command("subjects"))
async def cmd_subjects(message: Message) -> None:
    if not data["subjects"]:
        await message.answer("Пока нет настроенных предметов.")
        return
    lines = [f"{emoji} {name}" for name, emoji in data["subjects"].items()]
    await message.answer("\n".join(lines))


@router.message(Command("teachers"))
async def cmd_teachers(message: Message) -> None:
    if not data["teachers"]:
        await message.answer("Пока нет настроенных преподавателей.")
        return
    lines = [f"{emoji} {name}" for name, emoji in data["teachers"].items()]
    await message.answer("\n".join(lines))


@router.message(F.document)
async def handle_document(message: Message) -> None:
    doc = message.document
    file_name = (doc.file_name or "").lower()
    if not file_name.endswith(".docx"):
        await message.answer("Пришли файл в формате .docx с таблицей расписания.")
        return

    caption = (message.caption or "").strip()
    dm = DATE_IN_CAPTION_RE.search(caption)
    if dm:
        dd, mm = int(dm.group(1)), int(dm.group(2))
    else:
        now = datetime.now(KYIV_TZ)
        dd, mm = now.day, now.month
    target_date_str = f"{dd:02d}.{mm:02d}"

    group_name = data.get("group", DEFAULT_GROUP)

    try:
        tg_file = await message.bot.get_file(doc.file_id)
        buf = await message.bot.download_file(tg_file.file_path)
        file_bytes = buf.read()
        entries = extract_group_schedule(file_bytes, group_name)
    except Exception:
        logger.exception("Не удалось разобрать .docx")
        await message.answer(
            "Не получилось прочитать файл. Проверь, что это .docx с таблицей "
            "расписания (как выгружает деканат)."
        )
        return

    if not entries:
        await message.answer(f"Не нашёл группу «{group_name}» в этом файле.")
        return

    picked = pick_for_date(entries, target_date_str)
    blocks = []
    uncertain_labels = []
    for pair_num in sorted(picked):
        for subgroup, entry, note in picked[pair_num]:
            blocks.append(format_docx_entry(pair_num, entry, subgroup))
            if note == "uncertain":
                label = str(pair_num)
                if subgroup is not None:
                    label += f" ({SUBGROUP_LABEL.get(subgroup, subgroup)})"
                uncertain_labels.append(label)

    if not blocks:
        await message.answer(f"Не нашёл ни одной пары для группы «{group_name}» в файле.")
        return

    text = f"Розклад для {group_name} на {target_date_str}:\n\n" + "\n\n".join(blocks)
    if uncertain_labels:
        pairs_str = ", ".join(uncertain_labels)
        text += (
            f"\n\n⚠️ Для пари(-ар) {pairs_str} не знайшов точну дату в файлі — "
            "показав типовий запис, звір вручну."
        )

    try:
        await message.answer(text, parse_mode=ParseMode.MARKDOWN, disable_web_page_preview=True)
    except TelegramBadRequest:
        await message.answer(text, disable_web_page_preview=True)


@router.message(F.text)
async def handle_text(message: Message) -> None:
    result = format_schedule(message.text or "")
    if not result:
        await message.answer(
            "Не нашёл в тексте ни одной пары в формате «N пара». "
            "Пришли расписание как в примере, файл .docx с расписанием, или посмотри /help."
        )
        return
    try:
        await message.answer(result, parse_mode=ParseMode.MARKDOWN, disable_web_page_preview=True)
    except TelegramBadRequest:
        # на случай спецсимволов в названии предмета/преподавателя, ломающих markdown
        await message.answer(result, disable_web_page_preview=True)


async def main() -> None:
    token = os.environ.get("BOT_TOKEN")
    if not token:
        token_file = BASE_DIR / "token.txt"
        if token_file.exists():
            token = token_file.read_text(encoding="utf-8").strip()
    if not token:
        raise SystemExit(
            "Не найден токен бота. Установи переменную окружения BOT_TOKEN "
            "или создай рядом со скриптом файл token.txt с токеном внутри."
        )

    bot = Bot(token=token)
    dp = Dispatcher()
    dp.include_router(router)

    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
